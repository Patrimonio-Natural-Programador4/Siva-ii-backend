import shutil
import subprocess
import unicodedata
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

_PARAMETROS_DESEADOS = [
    "RUBRO",
    "No. PROCESO",
    "LUGAR DE EJECUCIÓN",
]


def extraer_parametros(nombre_archivo: str, contenido: bytes) -> list[dict[str, str]]:
    extension = Path(nombre_archivo or "").suffix.lower()
    if extension not in {".doc", ".docx"}:
        raise ValueError("El archivo debe tener extensión .doc o .docx")

    with TemporaryDirectory() as directorio_temporal:
        archivo_original = Path(directorio_temporal, f"documento{extension}")
        archivo_original.write_bytes(contenido)
        archivo_docx = _convertir_a_docx(archivo_original)
        return _extraer_tablas_parametros(archivo_docx)


def _convertir_a_docx(archivo: Path) -> Path:
    if archivo.suffix.lower() == ".docx":
        return archivo

    libreoffice = shutil.which("soffice") or shutil.which("libreoffice")
    if libreoffice is None:
        raise ValueError(
            "La lectura de archivos .doc requiere LibreOffice. "
            "Instale libreoffice-writer en el servidor o cargue el documento como .docx."
        )

    archivo_convertido = archivo.with_suffix(".docx")
    try:
        subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation={archivo.parent.joinpath('perfil-libreoffice').as_uri()}",
                "--convert-to",
                "docx",
                "--outdir",
                str(archivo.parent),
                str(archivo),
            ],
            check=True,
            capture_output=True,
            text=True,
            timeout=60,
        )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
        raise ValueError(f"No fue posible convertir el archivo .doc: {error}") from error

    if not archivo_convertido.exists():
        raise ValueError("LibreOffice no generó el archivo .docx convertido")

    return archivo_convertido


def _extraer_tablas_parametros(archivo: Path) -> list[dict[str, str]]:
    documento = Document(str(archivo))
    parametros: list[dict[str, str]] = []

    for nombre in _PARAMETROS_DESEADOS:
        valor = _buscar_valor_en_tablas(documento.tables, nombre)
        if not valor:
            valor = _buscar_valor_en_parrafos(documento, nombre)
        if not valor:
            continue
        parametros.append({"nombre": nombre, "valor": valor})
    return parametros


def _iterar_tablas(tablas):
    for tabla in tablas:
        yield tabla
        for fila in tabla.rows:
            for celda in fila.cells:
                yield from _iterar_tablas(celda.tables)


def _buscar_valor_en_tablas(tablas, nombre_buscado: str) -> str:
    nombre_normalizado = _normalizar_nombre(nombre_buscado)

    for tabla in _iterar_tablas(tablas):
        fila_inicial = 1 if _es_tabla_parametros(tabla) else 0
        if fila_inicial == 0 and not _es_tabla_pares_parametros(tabla):
            continue

        for fila in tabla.rows[fila_inicial:]:
            nombre = _normalizar_texto(fila.cells[0].text)
            if _normalizar_nombre(nombre) == nombre_normalizado:
                return _normalizar_texto(fila.cells[1].text)

    return ""


def _buscar_valor_en_parrafos(documento: Document, nombre_buscado: str) -> str:
    nombre_normalizado = _normalizar_nombre(nombre_buscado)
    parrafos = [_normalizar_texto(parrafo.text) for parrafo in documento.paragraphs]

    for indice, texto in enumerate(parrafos):
        nombre = texto.rstrip(".:")
        if _normalizar_nombre(nombre) != nombre_normalizado:
            continue

        return _obtener_valor_siguiente(parrafos, indice + 1)

    return ""


def _obtener_valor_siguiente(parrafos: list[str], indice_inicial: int) -> str:
    valores: list[str] = []
    for texto in parrafos[indice_inicial:]:
        if not texto:
            if valores:
                break
            continue
        valores.append(texto)

    return " ".join(valores)


def _es_tabla_parametros(tabla) -> bool:
    if not tabla.rows or len(tabla.rows[0].cells) < 2:
        return False

    encabezados = [_normalizar_texto(celda.text).upper() for celda in tabla.rows[0].cells[:2]]
    return encabezados == ["NOMBRE", "VALOR"]


def _es_tabla_pares_parametros(tabla) -> bool:
    return bool(tabla.rows and all(len(fila.cells) == 2 for fila in tabla.rows))


def _normalizar_texto(texto: str) -> str:
    return " ".join(texto.split())


def _normalizar_nombre(nombre: str) -> str:
    sin_acentos = unicodedata.normalize("NFD", nombre)
    return "".join(caracter for caracter in sin_acentos if unicodedata.category(caracter) != "Mn").upper()