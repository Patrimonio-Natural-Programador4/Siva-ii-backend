from pathlib import Path
from html import escape
import re
from decimal import Decimal

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from config import STORAGE_PATH
from services.TdrDocumentService import PATRON_HTML, html_a_texto, reemplazar_en_parrafo


def construir_objetivos_actividades_productos(productos) -> str:
    hijos = {}
    for producto in productos:
        hijos.setdefault(producto.id_product, []).append(producto)

    def clave_orden(producto):
        codigo = str(producto.code or "")
        partes = tuple(
            (0, int(parte)) if parte.isdigit() else (1, parte.casefold())
            for parte in re.split(r"(\d+)", codigo)
        )
        return (not bool(codigo), partes, producto.id)

    def parrafo(etiqueta, producto):
        descripcion = escape(html_a_texto(producto.product or "")).replace("\n", "<br>")
        return f"<p><strong>{etiqueta}</strong> {descripcion}</p>"

    bloques = []
    objetivos = sorted(hijos.get(None, []) + hijos.get(0, []), key=clave_orden)
    for numero_objetivo, objetivo in enumerate(objetivos, start=1):
        bloques.append(parrafo(f"Objetivo espec\u00edfico {numero_objetivo}.", objetivo))
        numero_producto = 0
        actividades = sorted(hijos.get(objetivo.id, []), key=clave_orden)
        for numero_actividad, actividad in enumerate(actividades, start=1):
            bloques.append(parrafo(f"Actividad {numero_objetivo}.{numero_actividad}.", actividad))
            for producto in sorted(hijos.get(actividad.id, []), key=clave_orden):
                numero_producto += 1
                bloques.append(parrafo(f"Producto {numero_producto}.", producto))
    return "".join(bloques)


def insertar_tabla_desembolsos(documento, parrafo, desembolsos) -> None:
    grupos = {}
    for desembolso, producto in desembolsos:
        grupo = grupos.setdefault(desembolso.id, {"desembolso": desembolso, "productos": {}})
        grupo["productos"].setdefault(producto.id, producto)

    tabla = documento.add_table(rows=1, cols=2)
    tabla.autofit = False
    seccion = documento.sections[0]
    ancho = seccion.page_width - seccion.left_margin - seccion.right_margin
    anchos = (int(ancho * 0.77), int(ancho * 0.23))
    for columna, ancho_columna in zip(tabla.columns, anchos):
        columna.width = ancho_columna
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), "4")
        borde.set(qn("w:color"), "000000")
        bordes.append(borde)
    tabla._tbl.tblPr.append(bordes)
    for celda, titulo in zip(tabla.rows[0].cells, ("Productos", "MONTO ($)")):
        celda.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        celda.paragraphs[0].add_run(titulo).bold = True
        fondo = OxmlElement("w:shd")
        fondo.set(qn("w:fill"), "D9D9D9")
        celda._tc.get_or_add_tcPr().append(fondo)
    repetir = OxmlElement("w:tblHeader")
    tabla.rows[0]._tr.get_or_add_trPr().append(repetir)

    def moneda(valor):
        return "$ " + f"{valor:,.2f}".replace(",", "#").replace(".", ",").replace("#", ".")

    total = Decimal("0")
    grupos_ordenados = sorted(grupos.values(), key=lambda grupo: (
        grupo["desembolso"].number_disbursement is None,
        grupo["desembolso"].number_disbursement or 0,
        grupo["desembolso"].id,
    ))
    ordinales = {1: "PRIMER", 2: "SEGUNDO", 3: "TERCER", 4: "CUARTO", 5: "QUINTO", 6: "SEXTO", 7: "SÉPTIMO", 8: "OCTAVO", 9: "NOVENO", 10: "DÉCIMO"}
    for indice, grupo in enumerate(grupos_ordenados, start=1):
        desembolso = grupo["desembolso"]
        numero = desembolso.number_disbursement
        numero = numero if numero is not None else indice
        productos_celda, monto_celda = tabla.add_row().cells
        titulo = productos_celda.paragraphs[0]
        titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        titulo.add_run(f"{ordinales.get(numero, str(numero))} DESEMBOLSO").bold = True
        for producto in grupo["productos"].values():
            texto = productos_celda.add_paragraph()
            texto.add_run(f"Producto {producto.code if producto.code is not None else producto.id}: ").bold = True
            texto.add_run(html_a_texto(producto.product or ""))
        valor = Decimal(str(desembolso.value)) if desembolso.value is not None else Decimal("0")
        total += valor
        monto_celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        monto_celda.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        monto_celda.paragraphs[0].add_run(moneda(valor)).bold = True
    if not grupos:
        tabla.add_row().cells[0].text = "Sin desembolsos asociados"
    etiqueta_total, monto_total = tabla.add_row().cells
    etiqueta_total.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    etiqueta_total.paragraphs[0].add_run("TOTAL").bold = True
    monto_total.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    monto_total.paragraphs[0].add_run(moneda(total)).bold = True
    for fila in tabla.rows:
        for celda, ancho_celda in zip(fila.cells, anchos):
            celda.width = ancho_celda
    parrafo._p.addprevious(tabla._tbl)
    parrafo._p.getparent().remove(parrafo._p)


def generar_formato_ep(diccionario: dict[str, object], code: str) -> Path:
    base = Path(STORAGE_PATH or Path(__file__).resolve().parents[1])
    directorio = base / "soportes" / "acuerdos" / "estudios_previos"
    plantilla = directorio / "template" / "plantilla_estudios_previos.docx"
    documento = Document(str(plantilla))
    valores = {
        marcador.removeprefix("<<").removesuffix(">>"): html_a_texto(
           str(valor) if valor is not None else ""
        )
        for marcador, valor in diccionario.items()
        if marcador != "<<DESEMBOLSOS>>"
    }
    valores_html = {
        marcador.removeprefix("<<").removesuffix(">>"): valor
        for marcador, valor in diccionario.items()
        if isinstance(valor, str) and PATRON_HTML.search(valor)
    }
    partes = [documento.part]
    for seccion in documento.sections:
        for contenedor in (
            seccion.header, seccion.first_page_header, seccion.even_page_header,
            seccion.footer, seccion.first_page_footer, seccion.even_page_footer,
        ):
            if not contenedor.is_linked_to_previous and contenedor.part not in partes:
                partes.append(contenedor.part)
    for parte in partes:
        for elemento in list(parte.element.xpath(".//w:p")):
            parrafo = Paragraph(elemento, parte)
            if "<<DESEMBOLSOS>>" in parrafo.text and "<<DESEMBOLSOS>>" in diccionario:
                if parrafo.text.strip() != "<<DESEMBOLSOS>>":
                    raise ValueError("El marcador <<DESEMBOLSOS>> debe estar en un parrafo independiente")
                insertar_tabla_desembolsos(documento, parrafo, diccionario["<<DESEMBOLSOS>>"])
                continue
            for fragmento in parrafo.runs:
                if "<<ALCANCE>>" in fragmento.text and "<<ALCANCE>>" not in fragmento.text:
                    fragmento.text = fragmento.text.replace("<<ALCANCE>>", "<<ALCANCE>>")
            reemplazar_en_parrafo(parrafo, valores, valores_html)

    codigo_archivo = "".join(
        caracter if caracter.isalnum() or caracter in "-_" else "_"
        for caracter in str(code)
    )
    salida = directorio / "formato" / f"ep_{codigo_archivo}.docx"
    salida.parent.mkdir(parents=True, exist_ok=True)
    documento.save(str(salida))
    return salida

    