import os
import tempfile
import openpyxl
from openpyxl.styles import Font
from sqlalchemy.orm import Session
from repository.TravelExportRepository import (
    get_all_general,
    get_all_itineraries,
    get_all_accommodations,
    get_all_legalizations
)
from entity.travel_requests import TravelRequests
from entity.travel_requests import TravelRequests

def auto_adjust_columns(sheet):
    for col in sheet.columns:
        max_length = 0
        column = col[0].column_letter # trae el nombre de la columna
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        sheet.column_dimensions[column].width = adjusted_width

def build_excel_export(db: Session) -> str:
    wb = openpyxl.Workbook()
    
    # 1. LISTADO CONSOLIDADO
    ws_general = wb.active
    ws_general.title = "LISTADO CONSOLIDADO"
    general_data = get_all_general(db)
    leg_data = get_all_legalizations(db)
    
    req_data = db.query(TravelRequests.code, TravelRequests.request_date, TravelRequests.created_at).all()
    req_dates = {}
    for r in req_data:
        req_dates[r.code] = r.request_date or (r.created_at.date() if r.created_at else None)
        
    leg_totals = {}
    for leg in leg_data:
        leg_totals[leg.code] = leg_totals.get(leg.code, 0) + (leg.total_paid or 0)
    
    headers_gen = [
        "Codigo de viaje", "Usuario", "Fecha de solicitud", "Fecha y hora de salida", "Fecha y hora de llegada",
        "Requiere anticipo", "Monto solicitado", "Total días", "Estado de la solicitud",
        "Rubro", "Categoría Gasto", "Objetivo de la actividad", "Observaciones adicionales",
        "Supervisor", "Programa", "Es un viaje internacional", "El viaje es para un invitado",
        "Fecha de nacimiento", "Celular", "Parentesco", "Telefono",
        "Contacto de emergencia", "Incluye alimentación", "Valor Factura"
    ]
    ws_general.append(headers_gen)
    for cell in ws_general[1]:
        cell.font = Font(bold=True)
        
    for item in general_data:
        valor_factura = leg_totals.get(item.code, 0)
        fecha_solicitud = req_dates.get(item.code, item.request_date)
        
        ws_general.append([
            item.code, item.traveler, fecha_solicitud, item.travel_start_date, item.travel_end_date,
            "Sí" if item.requires_advance_payment else "No", item.requested_advance_amount,
            item.total_days, item.status, item.budget_item, item.expense_category,
            item.travel_purpose, item.general_additional_comments, item.supervisor,
            item.program, "Sí" if item.is_international_travel else "No",
            "Sí" if item.is_guest_travel else "No",
            item.birth_date, item.mobile_phone, item.emergency_relationship,
            item.emergency_contact_phone, item.emergency_contact_name,
            "Sí" if item.includes_food else "No", valor_factura
        ])
    auto_adjust_columns(ws_general)

    # 2. ITINERARIO
    ws_itinerary = wb.create_sheet("ITINERARIO")
    itinerary_data = get_all_itineraries(db)
    headers_itin = [
        "Codigo de viaje", "Fecha de solicitud", "Usuario", "Departamento origen", "Municipio origen", "Departamento destino", "Municipio destino",
        "Requiere tiquetes", "Fecha - Hora", "Hora", "Destino zona rural", "Observaciones"
    ]
    ws_itinerary.append(headers_itin)
    for cell in ws_itinerary[1]:
        cell.font = Font(bold=True)
        
    for item in itinerary_data:
        fecha_solicitud = req_dates.get(item.code, "")
        ws_itinerary.append([
            item.code, fecha_solicitud, item.traveler, item.origin_department, item.origin_municipality,
            item.destination_department, item.destination_municipality,
            "Sí" if item.requires_air_tickets else "No", item.travel_date,
            item.estimated_departure_time, "Sí" if item.is_rural_destination else "No",
            item.itinerary_additional_comments
        ])
    auto_adjust_columns(ws_itinerary)

    # 3. HOSPEDAJE
    ws_accom = wb.create_sheet("HOSPEDAJE")
    accom_data = get_all_accommodations(db)
    headers_acc = [
        "Codigo de viaje", "Fecha de solicitud", "Usuario", "Departamento", "Ciudad", "Fecha llegada",
        "Fecha salida", "Observaciones"
    ]
    ws_accom.append(headers_acc)
    for cell in ws_accom[1]:
        cell.font = Font(bold=True)
        
    for item in accom_data:
        fecha_solicitud = req_dates.get(item.code, "")
        ws_accom.append([
            item.code, fecha_solicitud, item.traveler, item.accommodation_department,
            item.accommodation_municipality, item.check_in_date, item.check_out_date,
            item.accommodation_additional_comments
        ])
    auto_adjust_columns(ws_accom)

    # 4. LEGALIZACIONES
    ws_leg = wb.create_sheet("LEGALIZACIONES")
    leg_data = get_all_legalizations(db)
    headers_leg = [
        "Codigo de viaje", "Fecha de solicitud", "Usuario", "Fecha factura", "Número documento", "Beneficiario",
        "NIT / CC", "Régimen", "Subtotal", "IVA",
        "(%) Retención", "Retención", "Valor Cancelado", "Observaciones de desembolso",
        "Observaciones"
    ]
    ws_leg.append(headers_leg)
    for cell in ws_leg[1]:
        cell.font = Font(bold=True)
        
    for item in leg_data:
        fecha_solicitud = req_dates.get(item.code, "")
        ws_leg.append([
            item.code, fecha_solicitud, item.traveler, item.legalization_date, item.transaction_number,
            item.legalization_beneficiary, item.beneficiary_nit, item.regimen_type,
            item.legalized_subtotal, item.legalized_iva, item.retention_percentage,
            item.retention, item.total_paid, item.outlay_observations, item.legalization_observations
        ])
    auto_adjust_columns(ws_leg)

    fd, temp_path = tempfile.mkstemp(suffix=".xlsx", prefix="export_viajes_")
    os.close(fd)
    wb.save(temp_path)
    return temp_path
